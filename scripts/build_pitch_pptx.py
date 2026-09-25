"""Generate BayouCare DevDays pitch deck as PowerPoint."""

from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(ROOT, "BayouCare-DevDays-Pitch.pptx")
OUT_ALT = os.path.join(ROOT, "BayouCare-DevDays-Pitch-v2.pptx")
LOGO = os.path.join(ROOT, "public", "bayoucare-logo.png")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

G900 = RGBColor(0x06, 0x28, 0x1C)
G800 = RGBColor(0x0A, 0x3D, 0x2B)
G700 = RGBColor(0x0F, 0x6B, 0x4C)
G600 = RGBColor(0x14, 0x95, 0x67)
G100 = RGBColor(0xD8, 0xF5, 0xE8)
G50 = RGBColor(0xEE, 0xFB, 0xF4)
INK = RGBColor(0x14, 0x24, 0x1C)
SOFT = RGBColor(0x51, 0x63, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF7, 0xFA, 0xF8)
CORAL = RGBColor(0xE0, 0x5A, 0x3C)
LINE = RGBColor(0xD5, 0xE0, 0xD9)
FRAG_BG = RGBColor(0xFB, 0xE5, 0xDD)
FRAG_LINE = RGBColor(0xE8, 0xB0, 0xA0)
FRAG_INK = RGBColor(0x7A, 0x34, 0x24)
MUTED_DOT = RGBColor(0xC5, 0xD5, 0xCC)
DARK_DOT = RGBColor(0x4A, 0x7A, 0x66)
CHIP_LINE = RGBColor(0xC5, 0xE5, 0xD4)
TEAL_SOFT = RGBColor(0x8F, 0xC9, 0xAD)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text(shape, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return tf


def add_para(tf, text, size=16, bold=False, color=INK, space_before=6, align=PP_ALIGN.LEFT, font="Calibri"):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return p


def fill_solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rounded_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_solid(shape, fill)
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.12
    except Exception:
        pass
    return shape


def blank_slide():
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color=CREAM):
    fill_solid(
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height),
        color,
    )


def footer(slide, text, n, total=9, dark=False):
    color = TEAL_SOFT if dark else SOFT
    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(8), Inches(0.3))
    add_text(box, text, size=11, bold=False, color=color)
    x0 = Inches(10.15)
    for i in range(total):
        d = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            x0 + Inches(i * 0.28),
            Inches(7.12),
            Inches(0.14) if i + 1 != n else Inches(0.28),
            Inches(0.14),
        )
        fill_solid(d, G600 if i + 1 == n else (DARK_DOT if dark else MUTED_DOT))


# ---- Slide 1 ----
s = blank_slide()
bg(s, G800)
accent = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(-1.5), Inches(7), Inches(7))
fill_solid(accent, G700)

if os.path.exists(LOGO):
    s.shapes.add_picture(LOGO, Inches(0.9), Inches(1.6), Inches(1.15), Inches(1.15))

k = s.shapes.add_textbox(Inches(2.3), Inches(1.85), Inches(9), Inches(0.4))
add_text(k, "DEVDAYS LOUISIANA  ·  OCHSNER HEALTH CHALLENGE", size=12, bold=True, color=TEAL_SOFT)

title = s.shapes.add_textbox(Inches(0.9), Inches(3.0), Inches(11), Inches(1.2))
add_text(title, "BayouCare", size=60, bold=True, color=WHITE, font="Georgia")

sub = s.shapes.add_textbox(Inches(0.9), Inches(4.25), Inches(10), Inches(1))
tf = add_text(sub, "Making cancer care easier to navigate —", size=24, color=G100)
add_para(tf, "wherever patients live.", size=24, color=G100, space_before=2)

badges = ["Patient-centered", "Louisiana-first", "Clinic + population"]
x = Inches(0.9)
for b in badges:
    w = Inches(2.35)
    rounded_rect(s, x, Inches(5.55), w, Inches(0.42), G700)
    tb = s.shapes.add_textbox(x, Inches(5.58), w, Inches(0.36))
    add_text(tb, b, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    x += Inches(2.55)

footer(s, "Cancer care navigation", 1, dark=True)

# ---- Slide 2 ----
s = blank_slide()
bg(s, CREAM)
k = s.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(6), Inches(0.35))
add_text(k, "THE PROBLEM", size=12, bold=True, color=G700)
t = s.shapes.add_textbox(Inches(0.7), Inches(0.75), Inches(8), Inches(0.7))
add_text(t, "Meet Darlene.", size=36, bold=True, color=G900, font="Georgia")

rounded_rect(s, Inches(0.7), Inches(1.7), Inches(6.6), Inches(4.35), WHITE, LINE)
name = s.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(6), Inches(0.45))
add_text(name, "Darlene Fontenot", size=26, bold=True, color=G900, font="Georgia")
meta = s.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(6), Inches(0.55))
add_text(
    meta,
    "Alexandria, LA  ·  Stage II breast cancer  ·  day 4 of chemo  ·  42 miles from clinic",
    size=13,
    color=SOFT,
)
fever = s.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(0.9), Inches(0.35))
add_text(fever, "2:14 AM", size=14, bold=True, color=CORAL)
fever2 = s.shapes.add_textbox(Inches(1.9), Inches(3.0), Inches(5), Inches(0.35))
add_text(fever2, "— a fever starts.", size=14, color=INK)

qs = [
    "Is this dangerous?",
    "Should she go to the ER?",
    "Who should she contact?",
    "How will she get there?",
    "Does her care team know?",
]
y = Inches(3.45)
for q in qs:
    rounded_rect(s, Inches(1.0), y, Inches(5.9), Inches(0.38), G50, CHIP_LINE)
    tb = s.shapes.add_textbox(Inches(1.15), y + Inches(0.02), Inches(5.6), Inches(0.34))
    add_text(tb, q, size=13, bold=True, color=G800)
    y += Inches(0.45)

frags = [
    "Symptoms",
    "Transportation",
    "Appointments",
    "Caregivers",
    "Clinics",
    "Resources",
    "Treatment plans",
    "Between-visit gaps",
]
x0, y0 = Inches(7.7), Inches(1.7)
for i, f in enumerate(frags):
    col, row = i % 2, i // 2
    left = x0 + Inches(col * 2.55)
    top = y0 + Inches(row * 0.7)
    rounded_rect(s, left, top, Inches(2.4), Inches(0.55), FRAG_BG, FRAG_LINE)
    tb = s.shapes.add_textbox(left, top + Inches(0.08), Inches(2.4), Inches(0.4))
    add_text(tb, f, size=12, bold=True, color=FRAG_INK, align=PP_ALIGN.CENTER)

rounded_rect(s, Inches(0.7), Inches(6.2), Inches(11.9), Inches(0.7), G50)
tb = s.shapes.add_textbox(Inches(0.95), Inches(6.32), Inches(11.4), Inches(0.5))
add_text(
    tb,
    "The patient is often responsible for connecting a healthcare system that was never designed to feel connected.",
    size=14,
    bold=True,
    color=G900,
)
footer(s, "Patient perspective first", 2)

# ---- Slide 3: Problem → How BayouCare solves it ----
s = blank_slide()
bg(s, CREAM)
k = s.shapes.add_textbox(Inches(0.7), Inches(0.3), Inches(8), Inches(0.3))
add_text(k, "PROBLEM → SOLUTION", size=12, bold=True, color=G700)
t = s.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(12), Inches(0.55))
add_text(t, "What gets in the way — and how BayouCare helps.", size=28, bold=True, color=G900, font="Georgia")

# Column headers
rounded_rect(s, Inches(0.55), Inches(1.25), Inches(5.7), Inches(0.45), FRAG_BG, FRAG_LINE)
ht = s.shapes.add_textbox(Inches(0.7), Inches(1.32), Inches(5.4), Inches(0.35))
add_text(ht, "The problem", size=14, bold=True, color=FRAG_INK)

rounded_rect(s, Inches(7.05), Inches(1.25), Inches(5.7), Inches(0.45), G100, CHIP_LINE)
ht2 = s.shapes.add_textbox(Inches(7.2), Inches(1.32), Inches(5.4), Inches(0.35))
add_text(ht2, "How BayouCare helps", size=14, bold=True, color=G800)

pairs = [
    (
        "Limited access — patients live far from clinics and specialists",
        "Bridges the gap between patients, caregivers, and clinicians in one phone-first journey",
    ),
    (
        "Fragmented information across apps, calls, and paperwork",
        "Keeps the care plan, check-ins, and team updates in one coordinated place",
    ),
    (
        "Concerning symptoms go unseen between visits",
        "Surfaces risk early so the care team can respond before an ER trip",
    ),
    (
        "Transportation and logistics block treatment",
        "Connects rides and support actions into the clinic workflow — not a separate chase",
    ),
    (
        "Families don't know how to help",
        "Gives caregivers a clear role: shared tasks, updates, and next steps",
    ),
    (
        "Clinics lack visibility into who needs help now",
        "Shows 7-day risk with a concrete next action — and parish-level access gaps",
    ),
]

y = Inches(1.85)
for problem, solve in pairs:
    rounded_rect(s, Inches(0.55), y, Inches(5.7), Inches(0.75), WHITE, LINE)
    pt = s.shapes.add_textbox(Inches(0.7), y + Inches(0.1), Inches(5.4), Inches(0.6))
    add_text(pt, problem, size=12, bold=True, color=INK)

    arrow = s.shapes.add_textbox(Inches(6.35), y + Inches(0.15), Inches(0.55), Inches(0.45))
    add_text(arrow, "→", size=20, bold=True, color=G600, align=PP_ALIGN.CENTER)

    rounded_rect(s, Inches(7.05), y, Inches(5.7), Inches(0.75), G50, CHIP_LINE)
    st = s.shapes.add_textbox(Inches(7.2), y + Inches(0.1), Inches(5.4), Inches(0.6))
    add_text(st, solve, size=12, bold=True, color=G800)
    y += Inches(0.82)

footer(s, "Gaps closed by one connected platform", 3)

# ---- Slide 4 ----
s = blank_slide()
bg(s, CREAM)
k = s.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(6), Inches(0.3))
add_text(k, "THE SOLUTION", size=12, bold=True, color=G700)
t = s.shapes.add_textbox(Inches(0.7), Inches(0.65), Inches(11), Inches(0.6))
add_text(t, "One platform. One care journey.", size=32, bold=True, color=G900, font="Georgia")
lead = s.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(11), Inches(0.4))
add_text(
    lead,
    "BayouCare connects the person receiving care with the people responsible for delivering it.",
    size=16,
    color=SOFT,
)

nodes = [
    ("Patient", ["Symptom check-ins", "Treatment journey", "Appointments & guidance", "Resources"], Inches(0.7), Inches(2.0)),
    ("Caregiver", ["Shared updates", "Task coordination", "Transportation help", "Support circle"], Inches(7.1), Inches(2.0)),
    ("Clinic", ["Risk visibility", "Intervention workflow", "Prior-auth & ops", "Between-visit monitoring"], Inches(0.7), Inches(4.55)),
    ("Population health", ["Parish-level insights", "Access gaps", "Outreach planning", "Geographic trends"], Inches(7.1), Inches(4.55)),
]
for title_n, items, left, top in nodes:
    rounded_rect(s, left, top, Inches(5.5), Inches(2.2), WHITE, LINE)
    tb = s.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5), Inches(0.35))
    add_text(tb, title_n, size=16, bold=True, color=G800)
    body = s.shapes.add_textbox(left + Inches(0.25), top + Inches(0.6), Inches(5), Inches(1.4))
    tf = body.text_frame
    tf.clear()
    for j, item in enumerate(items):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.space_before = Pt(2)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=13, color=SOFT)

rounded_rect(s, Inches(5.15), Inches(3.55), Inches(3.0), Inches(1.35), G800)
tb = s.shapes.add_textbox(Inches(5.15), Inches(3.75), Inches(3.0), Inches(0.5))
add_text(tb, "BayouCare", size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font="Georgia")
tb2 = s.shapes.add_textbox(Inches(5.15), Inches(4.25), Inches(3.0), Inches(0.35))
add_text(tb2, "care navigation", size=12, color=G100, align=PP_ALIGN.CENTER)
footer(s, "Not another dashboard — a connected journey", 4)

# ---- Slide 5 ----
s = blank_slide()
bg(s, CREAM)
k = s.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(6), Inches(0.3))
add_text(k, "HOW IT WORKS", size=12, bold=True, color=G700)
t = s.shapes.add_textbox(Inches(0.55), Inches(0.65), Inches(12), Inches(0.55))
add_text(t, "From a check-in to coordinated support.", size=30, bold=True, color=G900, font="Georgia")

steps = [
    ("1", "Patient checks in", "Symptoms and how they're feeling — in minutes, on a phone."),
    ("2", "Risk surfaces", "BayouCare flags concerning changes that need attention."),
    ("3", "Care team sees it", "Earlier visibility between visits — not after an ER trip."),
    ("4", "Support connects", "Rides, caregivers, clinics, and resources can move together."),
    ("5", "Patterns emerge", "Parish insights show where patients struggle geographically."),
]
for i, (num, title_s, body) in enumerate(steps):
    left = Inches(0.45 + i * 2.55)
    rounded_rect(s, left, Inches(1.7), Inches(2.4), Inches(3.6), WHITE, LINE)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.2), Inches(1.95), Inches(0.45), Inches(0.45))
    fill_solid(circ, G100)
    nt = s.shapes.add_textbox(left + Inches(0.2), Inches(1.98), Inches(0.45), Inches(0.4))
    add_text(nt, num, size=14, bold=True, color=G800, align=PP_ALIGN.CENTER)
    tt = s.shapes.add_textbox(left + Inches(0.2), Inches(2.6), Inches(2.0), Inches(0.9))
    add_text(tt, title_s, size=15, bold=True, color=G900)
    bt = s.shapes.add_textbox(left + Inches(0.2), Inches(3.55), Inches(2.0), Inches(1.5))
    add_text(bt, body, size=12, color=SOFT)
    if i < 4:
        ar = s.shapes.add_textbox(left + Inches(2.25), Inches(3.2), Inches(0.35), Inches(0.4))
        add_text(ar, "→", size=18, bold=True, color=G600)

rounded_rect(s, Inches(0.55), Inches(5.7), Inches(4.2), Inches(0.65), G800)
tb = s.shapes.add_textbox(Inches(0.55), Inches(5.8), Inches(4.2), Inches(0.45))
add_text(tb, "Let's see it in action  →", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
footer(s, "Story → live product", 5)

# ---- Slide 6 ----
s = blank_slide()
bg(s, CREAM)
k = s.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(6), Inches(0.3))
add_text(k, "PRODUCT", size=12, bold=True, color=G700)
t = s.shapes.add_textbox(Inches(0.7), Inches(0.65), Inches(12), Inches(0.55))
add_text(t, "From one patient to an entire population.", size=30, bold=True, color=G900, font="Georgia")

devices = [
    ("PATIENT", "My Care", "Check-ins, vitals, family help, Remi guidance — what matters today."),
    ("CLINIC", "Care Team + Ops", "7-day risk with a prescription, prior-auth, schedule risk — one model."),
    ("POPULATION", "64 parishes", "Unmet-need heat index — outreach follows burden, not city size."),
]
for i, (lab, title_d, blurb) in enumerate(devices):
    left = Inches(0.7 + i * 4.15)
    rounded_rect(s, left, Inches(1.55), Inches(3.9), Inches(3.7), WHITE, LINE)
    rounded_rect(s, left, Inches(1.55), Inches(3.9), Inches(0.45), G50, LINE)
    lt = s.shapes.add_textbox(left + Inches(0.25), Inches(2.2), Inches(3.4), Inches(0.3))
    add_text(lt, lab, size=11, bold=True, color=G700)
    tt = s.shapes.add_textbox(left + Inches(0.25), Inches(2.55), Inches(3.4), Inches(0.5))
    add_text(tt, title_d, size=22, bold=True, color=G900, font="Georgia")
    bt = s.shapes.add_textbox(left + Inches(0.25), Inches(3.2), Inches(3.4), Inches(1.2))
    add_text(bt, blurb, size=13, color=SOFT)
    rounded_rect(s, left + Inches(0.25), Inches(4.6), Inches(1.5), Inches(0.35), G100)
    pt = s.shapes.add_textbox(left + Inches(0.25), Inches(4.62), Inches(1.5), Inches(0.3))
    add_text(pt, "Live in demo", size=11, bold=True, color=G800, align=PP_ALIGN.CENTER)

rounded_rect(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(0.95), G800)
tb = s.shapes.add_textbox(Inches(1.0), Inches(5.7), Inches(10), Inches(0.35))
add_text(tb, "Switching to the live product", size=18, bold=True, color=WHITE)
tb2 = s.shapes.add_textbox(Inches(1.0), Inches(6.1), Inches(10.5), Inches(0.3))
add_text(
    tb2,
    "Working prototype — synthetic clinical data · real parish sources labeled on screen",
    size=12,
    color=G100,
)
ar = s.shapes.add_textbox(Inches(11.5), Inches(5.75), Inches(0.8), Inches(0.5))
add_text(ar, "→", size=28, bold=True, color=WHITE)
footer(s, "UI hero is the running app", 6)

# ---- Slide 7 ----
s = blank_slide()
bg(s, CREAM)
k = s.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(8), Inches(0.3))
add_text(k, "WHY BAYOUCARE IS DIFFERENT", size=12, bold=True, color=G700)
t = s.shapes.add_textbox(Inches(0.7), Inches(0.65), Inches(12), Inches(0.7))
add_text(t, "Most tools solve one part. We connect the journey.", size=28, bold=True, color=G900, font="Georgia")

levels = [
    ("INDIVIDUAL", "Patient navigation", "Symptom awareness, plain-language guidance, family support on a phone."),
    ("CARE TEAM", "Coordination & intervention", "Risk forecasts with actionable next steps, shared with clinic ops."),
    ("COMMUNITY", "Population insight", "Louisiana parish visibility so resources follow unmet need."),
]
for i, (tier, title_l, body) in enumerate(levels):
    left = Inches(0.7 + i * 4.15)
    rounded_rect(s, left, Inches(1.7), Inches(3.95), Inches(2.5), WHITE, LINE)
    tt = s.shapes.add_textbox(left + Inches(0.25), Inches(1.9), Inches(3.4), Inches(0.3))
    add_text(tt, tier, size=11, bold=True, color=G600)
    th = s.shapes.add_textbox(left + Inches(0.25), Inches(2.25), Inches(3.4), Inches(0.55))
    add_text(th, title_l, size=18, bold=True, color=G900, font="Georgia")
    bd = s.shapes.add_textbox(left + Inches(0.25), Inches(2.95), Inches(3.4), Inches(1.0))
    add_text(bd, body, size=13, color=SOFT)

diffs = [
    "Phone-first accessibility",
    "Patient + caregiver + clinic",
    "Rural logistics in-product",
    "Louisiana geographic context",
]
for i, d in enumerate(diffs):
    left = Inches(0.7 + i * 3.1)
    rounded_rect(s, left, Inches(4.5), Inches(2.95), Inches(0.7), G50, CHIP_LINE)
    tb = s.shapes.add_textbox(left + Inches(0.1), Inches(4.62), Inches(2.75), Inches(0.5))
    add_text(tb, d, size=12, bold=True, color=G800, align=PP_ALIGN.CENTER)

key = s.shapes.add_textbox(Inches(0.7), Inches(5.55), Inches(12), Inches(0.55))
add_text(key, "One care story across three altitudes of healthcare.", size=20, bold=True, color=G900, font="Georgia")
footer(s, "Connected, not bolted-on", 7)

# ---- Slide 8 ----
s = blank_slide()
bg(s, CREAM)
k = s.shapes.add_textbox(Inches(0.7), Inches(0.3), Inches(8), Inches(0.3))
add_text(k, "TECHNOLOGY", size=12, bold=True, color=G700)
t = s.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(12), Inches(0.55))
add_text(t, "Built like a real healthcare platform.", size=28, bold=True, color=G900, font="Georgia")

layers = [
    ("EXPERIENCE", "Patient app · Caregiver / family · Care Team · Clinic Ops · Population map"),
    ("APPLICATION", "Auth (Clerk + session) · Symptom check-ins · Risk / counterfactuals · Care coordination · Geospatial · Remi (DeepSeek)"),
    ("DATA", "Neon Postgres · Patient / clinic state · Parish / CDC sources · Synthetic demo cohort"),
    ("EXTERNAL", "MapLibre + Carto · ClinicalTrials.gov · DeepSeek API · Notifications / SMS demo"),
    ("INFRASTRUCTURE", "Vercel · Serverless APIs · Vite + React 19 · TypeScript"),
]
y = Inches(1.3)
for name, chips in layers:
    rounded_rect(s, Inches(0.55), y, Inches(8.3), Inches(0.85), WHITE, LINE)
    nt = s.shapes.add_textbox(Inches(0.75), y + Inches(0.08), Inches(7.9), Inches(0.25))
    add_text(nt, name, size=10, bold=True, color=G600)
    ct = s.shapes.add_textbox(Inches(0.75), y + Inches(0.35), Inches(7.9), Inches(0.4))
    add_text(ct, chips, size=12, color=G800)
    y += Inches(0.95)

rounded_rect(s, Inches(9.1), Inches(1.3), Inches(3.7), Inches(5.0), G800)
th = s.shapes.add_textbox(Inches(9.35), Inches(1.55), Inches(3.2), Inches(0.4))
add_text(th, "Working stack", size=18, bold=True, color=WHITE, font="Georgia")
desc = s.shapes.add_textbox(Inches(9.35), Inches(2.05), Inches(3.2), Inches(1.0))
add_text(
    desc,
    "Deployed end-to-end — not a clickable mock. Same risk engine drives Care Team and Clinic Ops.",
    size=12,
    color=G100,
)

rows = [
    ("Frontend", "React · Vite · Tailwind"),
    ("Auth", "Clerk"),
    ("Data", "Neon Postgres"),
    ("Maps", "MapLibre · Leaflet"),
    ("AI", "DeepSeek (Remi)"),
    ("Host", "Vercel Functions"),
]
yy = Inches(3.2)
for a, b in rows:
    line = s.shapes.add_textbox(Inches(9.35), yy, Inches(1.3), Inches(0.4))
    add_text(line, a, size=12, bold=True, color=WHITE)
    line2 = s.shapes.add_textbox(Inches(10.5), yy, Inches(2.1), Inches(0.4))
    add_text(line2, b, size=11, color=G100, align=PP_ALIGN.RIGHT)
    yy += Inches(0.42)

footer(s, "Confirmed project technologies only", 8)

# ---- Slide 9 ----
s = blank_slide()
bg(s, CREAM)
k = s.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(6), Inches(0.3))
add_text(k, "IMPACT", size=12, bold=True, color=G700)
t = s.shapes.add_textbox(Inches(0.7), Inches(0.65), Inches(12), Inches(0.7))
add_text(t, "More navigation. Earlier support. More good days.", size=28, bold=True, color=G900, font="Georgia")

outs = [
    ("For patients", "Less confusion navigating care."),
    ("For care teams", "Better visibility between visits."),
    ("For health systems", "Clearer view of where support is needed."),
]
for i, (h, b) in enumerate(outs):
    left = Inches(0.7 + i * 4.15)
    rounded_rect(s, left, Inches(1.7), Inches(3.95), Inches(1.8), WHITE, LINE)
    ht = s.shapes.add_textbox(left + Inches(0.25), Inches(1.95), Inches(3.4), Inches(0.35))
    add_text(ht, h, size=13, bold=True, color=G700)
    bt = s.shapes.add_textbox(left + Inches(0.25), Inches(2.4), Inches(3.4), Inches(0.8))
    add_text(bt, b, size=18, bold=True, color=G900, font="Georgia")

rounded_rect(s, Inches(0.7), Inches(3.8), Inches(11.9), Inches(1.45), G50, CHIP_LINE)
vt = s.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(11.3), Inches(0.55))
add_text(
    vt,
    "Start with Louisiana. Build a model that can scale anywhere geography makes healthcare harder to reach.",
    size=15,
    bold=True,
    color=G900,
)
vn = s.shapes.add_textbox(Inches(1.0), Inches(4.65), Inches(11.3), Inches(0.4))
add_text(
    vn,
    "Next step: pilot BayouCare with healthcare partners and real patient-navigation workflows.",
    size=13,
    color=SOFT,
)

close = s.shapes.add_textbox(Inches(0.7), Inches(5.55), Inches(12), Inches(0.5))
add_text(close, "Cancer is difficult enough. Navigating care shouldn't be.", size=22, bold=True, color=G900, font="Georgia")
brand = s.shapes.add_textbox(Inches(0.7), Inches(6.15), Inches(12), Inches(0.35))
add_text(brand, "BayouCare  ·  DevDays 2026", size=14, bold=True, color=G600)
footer(s, "Thank you", 9)

try:
    prs.save(OUT)
    saved = OUT
except PermissionError:
    prs.save(OUT_ALT)
    saved = OUT_ALT
print(saved)
print(f"slides={len(prs.slides)}")
